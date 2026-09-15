"""Endpoint tests for conversation transcript export.

The renderer itself is covered by ``test_docx_export.py``; these tests pin the
HTTP contract: auth, format validation, the Word content type, and that the
downloaded document really contains the conversation either way.
"""

import io
from datetime import UTC, datetime

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select

from app.core.config import Settings, get_settings
from app.core.security import get_current_user
from app.db.session import get_db
from app.main import app
from app.models.conversation import Conversation
from app.models.message import Message

pytestmark = pytest.mark.asyncio

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_TEST_SETTINGS = Settings(
    secret_key="test-secret-key-do-not-use-in-prod",
    database_url="sqlite+aiosqlite:///:memory:",
    access_token_expire_minutes=30,
)

OWNER = "test@example.com"  # seeded by conftest
OTHER = "other@example.com"


class _UserSwitch:
    def __init__(self, email: str = OWNER):
        self.email = email

    async def __call__(self):
        return {"email": self.email, "token_payload": {}}


def _install_overrides(db_session, switch: _UserSwitch):
    async def _override_db():
        yield db_session

    app.dependency_overrides[get_db] = _override_db
    app.dependency_overrides[get_settings] = lambda: _TEST_SETTINGS
    app.dependency_overrides[get_current_user] = switch


def _clear_overrides():
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_settings, None)
    app.dependency_overrides.pop(get_current_user, None)


def _client() -> AsyncClient:
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


async def _create_conversation(client: AsyncClient, title: str = "Design chat") -> dict:
    resp = await client.post("/api/v1/chat/conversations", json={"title": title})
    assert resp.status_code == 201, resp.text
    return resp.json()


async def _seed_messages(db_session, conversation_id: str, *messages: tuple[str, str]):
    for index, (role, content) in enumerate(messages):
        db_session.add(
            Message(
                id=f"export-{conversation_id[:8]}-{index}",
                conversation_id=conversation_id,
                role=role,
                content=content,
                created_at=datetime(2026, 1, 1, tzinfo=UTC),
                seq=index,
            )
        )
    await db_session.commit()


async def test_export_defaults_to_word_document(db_session):
    switch = _UserSwitch()
    _install_overrides(db_session, switch)
    try:
        async with _client() as c:
            conversation = await _create_conversation(c)
            await _seed_messages(
                db_session,
                conversation["id"],
                ("user", "How do I deploy?"),
                ("assistant", "## Steps\n\n1. Run `just deploy`\n2. Watch the logs"),
            )

            resp = await c.get(f"/api/v1/chat/conversations/{conversation['id']}/export")
        assert resp.status_code == 200
        assert resp.headers["content-type"] == DOCX_MEDIA_TYPE
        assert resp.content[:2] == b"PK"
        assert "attachment" in resp.headers["content-disposition"]
        assert "chat-design-chat.docx" in resp.headers["content-disposition"]

        document = Document(io.BytesIO(resp.content))
        text = "\n".join(p.text for p in document.paragraphs)
        assert "Design chat" in text
        assert OWNER in text
        assert "Assistant" in text
        assert "How do I deploy?" in text
        assert "Steps" in text
        assert "Run `just deploy`" not in text  # markdown, not literal backticks
        # The turn's own "## Steps" nested under the author heading, and its
        # numbered markdown list became numbered list paragraphs.
        assert ("Heading 3", "Steps") in [(p.style.name, p.text) for p in document.paragraphs]
        assert ("List Paragraph", "1. Run just deploy") in [
            (p.style.name, p.text) for p in document.paragraphs
        ]
    finally:
        _clear_overrides()


async def test_export_markdown_format(db_session):
    switch = _UserSwitch()
    _install_overrides(db_session, switch)
    try:
        async with _client() as c:
            conversation = await _create_conversation(c, title="Notes")
            await _seed_messages(db_session, conversation["id"], ("user", "hello"))
            resp = await c.get(
                f"/api/v1/chat/conversations/{conversation['id']}/export?format=markdown"
            )
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("text/markdown")
        assert "# Notes" in resp.text
        assert "hello" in resp.text
        # `?format=markdown` still downloads as .md — the API's format name is
        # not the file extension.
        assert resp.headers["content-disposition"] == 'attachment; filename="chat-notes.md"'
    finally:
        _clear_overrides()


async def test_export_omits_tool_messages_but_discloses_it(db_session):
    switch = _UserSwitch()
    _install_overrides(db_session, switch)
    try:
        async with _client() as c:
            conversation = await _create_conversation(c)
            await _seed_messages(
                db_session,
                conversation["id"],
                ("user", "search for me"),
                ("tool_call", '{"name": "web_search"}'),
                ("tool_result", '{"result": "..."}'),
                ("assistant", "Done."),
            )
            resp = await c.get(f"/api/v1/chat/conversations/{conversation['id']}/export")
        document = Document(io.BytesIO(resp.content))
        text = "\n".join(p.text for p in document.paragraphs)
        assert "web_search" not in text
        assert "2 tool messages omitted" in text
    finally:
        _clear_overrides()


async def test_export_primary_conversation_exports_only_its_session(db_session):
    """A topic switch bumps the epoch; the export follows the visible window."""
    switch = _UserSwitch()
    _install_overrides(db_session, switch)
    try:
        async with _client() as c:
            resp = await c.post("/api/v1/chat/conversations/primary", json={})
            assert resp.status_code in (200, 201), resp.text
            conversation_id = resp.json()["id"]

            db_session.add_all(
                [
                    Message(
                        id="epoch-old",
                        conversation_id=conversation_id,
                        role="user",
                        content="previous topic message",
                        created_at=datetime(2026, 1, 1, tzinfo=UTC),
                        seq=1,
                        session_epoch=0,
                    ),
                    Message(
                        id="epoch-new",
                        conversation_id=conversation_id,
                        role="user",
                        content="current topic message",
                        created_at=datetime(2026, 1, 1, tzinfo=UTC),
                        seq=2,
                        session_epoch=1,
                    ),
                ]
            )
            conversation = (
                await db_session.execute(
                    select(Conversation).where(Conversation.id == conversation_id)
                )
            ).scalar_one()
            conversation.session_epoch = 1
            conversation.title = "Primary"
            await db_session.commit()

            resp = await c.get(f"/api/v1/chat/conversations/{conversation_id}/export")
        assert resp.status_code == 200
        document = Document(io.BytesIO(resp.content))
        text = "\n".join(p.text for p in document.paragraphs)
        assert "current topic message" in text
        assert "previous topic message" not in text
    finally:
        _clear_overrides()


async def test_export_rejects_unknown_format(db_session):
    switch = _UserSwitch()
    _install_overrides(db_session, switch)
    try:
        async with _client() as c:
            conversation = await _create_conversation(c)
            resp = await c.get(f"/api/v1/chat/conversations/{conversation['id']}/export?format=pdf")
        assert resp.status_code == 422
    finally:
        _clear_overrides()


async def test_export_requires_ownership(db_session):
    switch = _UserSwitch()
    _install_overrides(db_session, switch)
    try:
        async with _client() as c:
            conversation = await _create_conversation(c)
            switch.email = OTHER
            resp = await c.get(f"/api/v1/chat/conversations/{conversation['id']}/export")
        assert resp.status_code == 404
    finally:
        _clear_overrides()


async def test_export_content_disposition_filename_is_safe(db_session):
    switch = _UserSwitch()
    _install_overrides(db_session, switch)
    try:
        async with _client() as c:
            conversation = await _create_conversation(c, title='Résumé / "draft" #2')
            resp = await c.get(f"/api/v1/chat/conversations/{conversation['id']}/export")
        disposition = resp.headers["content-disposition"]
        assert disposition == 'attachment; filename="chat-resume-draft-2.docx"'
    finally:
        _clear_overrides()

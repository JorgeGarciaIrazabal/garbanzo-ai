"""Unit tests for the Word/markdown transcript renderer.

These pin the *document structure* the feature request asks for: headings per
turn, real list paragraphs, code blocks, and tables — not literal markdown
punctuation leaked into the text.
"""

import io
from datetime import UTC, datetime

from docx import Document

from app.services.docx_export import (
    TranscriptSection,
    export_filename,
    render_transcript_docx,
    slugify,
)
from app.services.transcript_export import (
    build_export_footer,
    conversation_sections,
    render_transcript_markdown,
    room_sections,
)

MARKDOWN_BODY = """## Findings

A paragraph with **bold**, *italic*, `inline code`, ~~struck~~ and a [link](https://example.test).

- bullet one
- bullet two

1. first
2. second

- [ ] pending task
- [x] finished task

> quoted text

```python
def f():
    return 1
```

| Col A | Col B |
|-------|-------|
| 1     | 2     |

Final paragraph.
"""


def _document(data: bytes) -> Document:
    return Document(io.BytesIO(data))


def _rendered(body: str = MARKDOWN_BODY, **kwargs) -> Document:
    data = render_transcript_docx(
        title=kwargs.pop("title", "Design chat"),
        sections=[
            TranscriptSection(author="jorge@example.com", body=body, meta="2026-01-01 10:00")
        ],
        **kwargs,
    )
    # The bytes must be a real .docx package, not just any bytes.
    assert data[:2] == b"PK"
    return _document(data)


def _styles(document: Document) -> list[tuple[str, str]]:
    return [(p.style.name, p.text) for p in document.paragraphs if p.text.strip()]


class TestDocumentStructure:
    def test_title_and_author_headings(self):
        document = _rendered()
        assert document.paragraphs[0].text == "Design chat"
        assert document.paragraphs[0].style.name == "Title"
        styles = _styles(document)
        assert ("Heading 1", "jorge@example.com") in styles
        # A markdown "##" inside the turn nests beneath the author heading.
        assert ("Heading 3", "Findings") in styles

    def test_body_is_not_literal_markdown(self):
        document = _rendered()
        text = "\n".join(p.text for p in document.paragraphs)
        for punctuation in ("**", "~~", "```", "- [ ]", "| Col A |"):
            assert punctuation not in text

    def test_inline_formatting_becomes_runs(self):
        document = _rendered("Text with **bold**, *italic* and `code`.")
        paragraph = next(p for p in document.paragraphs if p.text.startswith("Text with"))
        runs = {run.text: run for run in paragraph.runs}
        assert runs["bold"].font.bold is True
        assert runs["italic"].font.italic is True
        assert runs["code"].font.name
        assert "code" in runs["code"].text

    def test_links_are_real_hyperlinks(self):
        document = _rendered("See [the docs](https://example.test/x).")
        paragraph = next(p for p in document.paragraphs if p.text.startswith("See"))
        hyperlinks = paragraph.hyperlinks
        assert [(h.text, h.url) for h in hyperlinks] == [("the docs", "https://example.test/x")]

    def test_bullets_and_nesting(self):
        document = _rendered("- one\n- two\n  - nested\n")
        styles = {text: style for style, text in _styles(document)}
        assert styles["one"] == "List Bullet"
        assert styles["two"] == "List Bullet"
        assert styles["nested"] == "List Bullet 2"

    def test_numbered_lists_keep_their_own_sequence(self):
        document = _rendered("1. one\n2. two\n\nText separates them.\n\n1. restarts\n")
        styles = _styles(document)
        assert ("List Paragraph", "1. one") in styles
        assert ("List Paragraph", "2. two") in styles
        # The second list restarts at 1 rather than continuing as 3.
        assert ("List Paragraph", "1. restarts") in styles

    def test_task_items_get_checkboxes(self):
        document = _rendered("- [ ] pending\n- [x] done\n")
        text = "\n".join(p.text for p in document.paragraphs)
        assert "☐ pending" in text
        assert "☑ done" in text
        assert "[ ]" not in text

    def test_code_block_is_monospace_and_keeps_indentation(self):
        document = _rendered("```python\nif x:\n    y = 1\n```\n")
        code_lines = [p for p in document.paragraphs if p.text and p.style.name == "No Spacing"]
        assert [p.text for p in code_lines] == ["if x:", "    y = 1"]
        assert code_lines[1].runs[0].font.name == "Consolas"

    def test_tables_become_word_tables(self):
        document = _rendered("| Col A | Col B |\n|-------|-------|\n| 1 | 2 |\n")
        assert len(document.tables) == 1
        table = document.tables[0]
        assert [[cell.text for cell in row.cells] for row in table.rows] == [
            ["Col A", "Col B"],
            ["1", "2"],
        ]

    def test_blank_body_produces_no_extra_content(self):
        document = _rendered("")
        styles = _styles(document)
        assert styles[0] == ("Title", "Design chat")
        assert ("Heading 1", "jorge@example.com") in styles
        # Title, caption, author heading, meta line — nothing else.
        assert len(styles) == 4

    def test_footer_and_omission_note(self):
        document = _rendered(footer=build_export_footer("Exported from Garbanzo AI", 3))
        text = "\n".join(p.text for p in document.paragraphs)
        assert "3 tool messages omitted" in text

    def test_malformed_markdown_never_loses_text(self):
        document = _rendered("Unclosed **bold and a | stray pipe\n\n<div>raw</div>\n")
        text = "\n".join(p.text for p in document.paragraphs)
        assert "Unclosed" in text
        assert "stray pipe" in text
        assert "raw" in text


class TestSectionBuilders:
    def test_conversation_sections_label_roles_and_count_machinery(self):
        class Message:
            def __init__(self, role, content):
                self.role = role
                self.content = content
                self.created_at = datetime(2026, 1, 1, tzinfo=UTC)

        sections, omitted = conversation_sections(
            [
                Message("user", "hi"),
                Message("assistant", "hello"),
                Message("tool_call", "{...}"),
                Message("tool_result", "{...}"),
                Message("assistant", "   "),
            ],
            user_label="jorge@example.com",
        )
        assert [(s.author, s.body) for s in sections] == [
            ("jorge@example.com", "hi"),
            ("Assistant", "hello"),
        ]
        assert omitted == 2

    def test_room_sections_resolve_agent_names_and_audio_notes(self):
        class Message:
            def __init__(self, **kwargs):
                self.role = kwargs.get("role", "user")
                self.content = kwargs.get("content", "")
                self.sender_user_id = kwargs.get("sender_user_id")
                self.sender_agent_id = kwargs.get("sender_agent_id")
                self.meta = kwargs.get("meta")
                self.created_at = datetime(2026, 1, 1, tzinfo=UTC)

        sections, omitted = room_sections(
            [
                Message(sender_user_id="jorge@example.com", content="hi"),
                Message(sender_agent_id="a1", content="hello"),
                Message(sender_agent_id="gone", content="unknown agent"),
                Message(
                    sender_user_id="jorge@example.com",
                    content="",
                    meta={"audio_note": {"duration_seconds": 75}},
                ),
                Message(role="tool_call", content="{...}"),
            ],
            {"a1": "Reviewer"},
        )
        assert [s.author for s in sections] == [
            "jorge@example.com",
            "🤖 Reviewer",
            "🤖 agent",
            "jorge@example.com",
        ]
        assert "Audio note (1:15)" in sections[3].body
        assert omitted == 1

    def test_markdown_export_has_a_heading_per_turn(self):
        markdown = render_transcript_markdown(
            title="Design chat",
            subtitle="Exported from Garbanzo AI",
            sections=[TranscriptSection(author="jorge", body="hi", meta="2026-01-01 10:00")],
        )
        assert markdown.startswith("# Design chat")
        assert "## jorge" in markdown
        assert "hi" in markdown


class TestFilenames:
    def test_slugify_drops_unsafe_characters(self):
        assert slugify("Room: Design chat / v2") == "room-design-chat-v2"
        assert slugify("¡Acentos y ñ!") == "acentos-y-n"
        assert slugify("") == "transcript"

    def test_export_filename_uses_prefix_and_extension(self):
        assert export_filename("chat", "Design chat", "docx") == "chat-design-chat.docx"
        assert export_filename("room", "!!!", "md") == "room-transcript.md"

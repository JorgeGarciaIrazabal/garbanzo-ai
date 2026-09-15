"""Render chat/room transcripts as Word (``.docx``) documents.

Assistant turns are markdown, so a transcript exported as plain text loses all
its structure. This module maps that markdown onto *native* Word constructs —
headings, bold/italic/code runs, real bullet and numbered-list paragraphs, code
blocks, and real tables — so the exported file opens in Word as an editable
document rather than a wall of literal ``**`` and ``-`` characters.

Used by both export endpoints (``GET /chat/conversations/{id}/export`` and
``GET /rooms/{id}/export``).
"""

from __future__ import annotations

import io
import re
import unicodedata
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import datetime

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from markdown_it import MarkdownIt
from markdown_it.token import Token

# python-docx ships no "Code" style; "No Spacing" + a monospace font is the
# closest built-in match and keeps code blocks compact and copyable.
MONOSPACE_FONT = "Consolas"
CODE_POINT_SIZE = Pt(9)
LINK_COLOR = RGBColor(0x0B, 0x57, 0xD0)
MUTED_COLOR = RGBColor(0x5F, 0x63, 0x68)

# Word only defines Heading 1..9 and List Bullet / 2 / 3.
MAX_HEADING_LEVEL = 9
BULLET_STYLES = ("List Bullet", "List Bullet 2", "List Bullet 3")

_TASK_ITEM_RE = re.compile(r"^\[([ xX])\]\s+")
_SLUG_RE = re.compile(r"[^a-z0-9]+")


@dataclass(frozen=True)
class TranscriptSection:
    """One turn: an author heading, an optional meta line, and a markdown body."""

    author: str
    body: str
    meta: str | None = None


def slugify(value: str, *, fallback: str = "transcript", max_length: int = 60) -> str:
    """A filesystem-safe ASCII token for export filenames."""
    normalized = unicodedata.normalize("NFKD", value or "")
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    slug = _SLUG_RE.sub("-", ascii_value.lower()).strip("-")
    return slug[:max_length].strip("-") or fallback


def export_filename(prefix: str, name: str, extension: str) -> str:
    return f"{prefix}-{slugify(name)}.{extension}"


def _markdown() -> MarkdownIt:
    # Same extensions the app's chat renderer enables (see lib/AGENTS.md:
    # tables / task lists / strikethrough all come from ExtensionSet.gitHubWeb).
    return MarkdownIt("commonmark").enable("table").enable("strikethrough")


class _ListFrame:
    """Numbering state for one (possibly nested) list."""

    def __init__(self, ordered: bool, start: int) -> None:
        self.ordered = ordered
        self.next_number = start

    def take(self) -> int | None:
        if not self.ordered:
            return None
        number = self.next_number
        self.next_number += 1
        return number


class _MarkdownWriter:
    """Streams markdown-it tokens into a python-docx document."""

    def __init__(self, document, *, heading_offset: int = 1) -> None:
        self.document = document
        self.heading_offset = heading_offset
        self.paragraph: Paragraph | None = None
        self.heading: Paragraph | None = None
        self.lists: list[_ListFrame] = []
        self.quote_depth = 0
        # List-item bookkeeping: the item's number is claimed on the first
        # paragraph it produces, and every later paragraph inside the same
        # item stays indented at the item's level without being numbered.
        self.item_open = False
        self.item_number: int | None = None
        self.item_numbered = False
        self.item_level = 0
        self.paragraph_in_item = False
        # Inline formatting marks.
        self.bold = 0
        self.italic = 0
        self.strike = 0
        self.link_stack: list[object] = []
        self.task_checkbox: str | None = None

    # ------------------------------------------------------------- block level

    def write(self, markdown_text: str) -> None:
        tokens = _markdown().parse(markdown_text or "")
        index = 0
        while index < len(tokens):
            token = tokens[index]
            if token.type == "table_open":
                index = self._consume_table(tokens, index)
                continue
            self._block(token)
            index += 1
        self._end_paragraph()

    def _block(self, token: Token) -> None:
        kind = token.type
        if kind == "heading_open":
            self._end_paragraph()
            self._start_heading(token.tag)
        elif kind == "heading_close":
            self._end_paragraph()
        elif kind == "inline":
            if self.heading is not None:
                paragraph, self.heading = self.heading, None
                self.paragraph = paragraph
            self._inline(token)
        elif kind == "paragraph_close":
            self._end_paragraph()
        elif kind in {"bullet_list_open", "ordered_list_open"}:
            self._end_paragraph()
            self.lists.append(_ListFrame(kind == "ordered_list_open", self._list_start(token)))
        elif kind in {"bullet_list_close", "ordered_list_close"}:
            self._end_paragraph()
            if self.lists:
                self.lists.pop()
        elif kind == "list_item_open":
            self._end_paragraph()
            self.item_open = True
            self.item_number = None
            frame = self.lists[-1] if self.lists else _ListFrame(False, 1)
            number = frame.take()
            if number is not None:
                self.item_number = number
                self.item_numbered = True
            self.item_level = max(len(self.lists) - 1, 0)
        elif kind == "list_item_close":
            self._end_paragraph()
            self.item_open = False
            self.item_number = None
            self.item_numbered = False
            self.task_checkbox = None
        elif kind == "blockquote_open":
            self._end_paragraph()
            self.quote_depth += 1
        elif kind == "blockquote_close":
            self._end_paragraph()
            self.quote_depth = max(self.quote_depth - 1, 0)
        elif kind in {"fence", "code_block"}:
            self._end_paragraph()
            self._add_code_block(token.content)
        elif kind == "hr":
            self._end_paragraph()
            self._add_separator()
        elif kind in {"html_block", "math_block"}:
            # Never drop author content: unsupported blocks keep their text.
            self._end_paragraph()
            self._text(token.content.strip())

    @staticmethod
    def _list_start(token: Token) -> int:
        try:
            return int(token.attrs.get("start", 1))
        except (TypeError, ValueError):
            return 1

    # -------------------------------------------------------- paragraph helpers

    def _end_paragraph(self) -> None:
        self.paragraph = None
        self.paragraph_in_item = False
        self.hyperlink = None
        self.task_checkbox = None

    def _ensure_paragraph(self) -> Paragraph:
        if self.paragraph is not None:
            return self.paragraph
        if self.item_open:
            self.paragraph_in_item = True
        style: str | None = None
        indent = 0.0
        hanging = False
        if self.paragraph_in_item:
            if self.item_numbered:
                # Numbered lists use a manual marker: the built-in
                # "List Number" style continues numbering across independent
                # lists in one document, which would make a transcript's
                # second list start where the first ended.
                style = "List Paragraph"
                indent = 0.25 * (self.item_level + 1)
                hanging = True
            else:
                style = BULLET_STYLES[min(self.item_level, len(BULLET_STYLES) - 1)]
                if self.item_level >= len(BULLET_STYLES):
                    indent = 0.25 * (self.item_level + 1)
        elif self.quote_depth:
            style = "Quote"

        paragraph = self.document.add_paragraph(style=style)
        if indent:
            paragraph.paragraph_format.left_indent = Inches(indent)
        if hanging:
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        self.paragraph = paragraph
        # A list item's marker rides on the paragraph that first materializes
        # it: the number for a numbered item and the checkbox for a task item.
        # Both are claimed once, so a following paragraph inside the same item
        # (a loose list) stays indented without repeating the marker.
        marker = ""
        if self.item_number is not None:
            marker += f"{self.item_number}. "
            self.item_number = None
        if self.task_checkbox:
            marker += f"{self.task_checkbox} "
            self.task_checkbox = None
        if marker:
            paragraph.add_run(marker)
        return paragraph

    def _start_heading(self, tag: str) -> None:
        try:
            level = int(tag[1:])
        except (IndexError, ValueError):
            level = 1
        level = max(1, min(level + self.heading_offset, MAX_HEADING_LEVEL))
        self.heading = self.document.add_heading(level=level)

    def _add_code_block(self, content: str) -> None:
        lines = content.rstrip("\n").split("\n") or [""]
        for line in lines:
            paragraph = self.document.add_paragraph(style="No Spacing")
            paragraph.paragraph_format.left_indent = Inches(0.25)
            run = paragraph.add_run(line)
            run.font.size = CODE_POINT_SIZE
            _force_font(run, MONOSPACE_FONT)

    def _add_separator(self) -> None:
        paragraph = self.document.add_paragraph()
        properties = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "BFBFBF")
        borders.append(bottom)
        properties.append(borders)

    # --------------------------------------------------------- inline content

    def _inline(self, token: Token) -> None:
        for child in token.children or []:
            self._inline_child(child)

    def _inline_child(self, token: Token) -> None:
        kind = token.type
        if kind == "text":
            self._text(token.content)
        elif kind in {"softbreak", "hardbreak"}:
            self._break()
        elif kind == "code_inline":
            run = self._add_run(token.content)
            run.font.size = CODE_POINT_SIZE
            _force_font(run, MONOSPACE_FONT)
        elif kind == "strong_open":
            self.bold += 1
        elif kind == "strong_close":
            self.bold = max(self.bold - 1, 0)
        elif kind == "em_open":
            self.italic += 1
        elif kind == "em_close":
            self.italic = max(self.italic - 1, 0)
        elif kind == "s_open":
            self.strike += 1
        elif kind == "s_close":
            self.strike = max(self.strike - 1, 0)
        elif kind == "link_open":
            self._open_link(str(token.attrs.get("href", "")))
        elif kind == "link_close":
            if self.link_stack:
                self.link_stack.pop()
        elif kind == "image":
            alt = token.attrGet("alt") or token.content or "image"
            self._text(f"[image: {alt}]")
        elif kind in {"html_inline", "math_inline"}:
            # Keep the raw text rather than silently dropping it.
            self._text(token.content)

    def _text(self, text: str) -> None:
        if self.task_checkbox is None and self.item_open:
            match = _TASK_ITEM_RE.match(text or "")
            if match:
                self.task_checkbox = "☑" if match.group(1).lower() == "x" else "☐"
                # _ensure_paragraph emits the checkbox with the item's marker
                # and clears it again, so no extra run is added here.
                self._ensure_paragraph()
                text = text[match.end() :]
        if not text:
            return
        run = self._add_run(text)
        if self.strike:
            run.font.strike = True

    def _break(self) -> None:
        self._add_run().add_break()

    def _open_link(self, href: str) -> None:
        paragraph = self._ensure_paragraph()
        element = OxmlElement("w:hyperlink")
        try:
            element.set(qn("r:id"), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
        except Exception:  # noqa: BLE001 - unusable URL: keep the label as plain text
            self.link_stack.append(None)
            return
        paragraph._p.append(element)
        self.link_stack.append(element)

    def _add_run(self, text: str = "") -> Run:
        paragraph = self._ensure_paragraph()
        run = paragraph.add_run(text)
        link = self.link_stack[-1] if self.link_stack else None
        if link is not None:
            link.append(run._r)
            run.font.underline = True
            run.font.color.rgb = LINK_COLOR
        elif self.link_stack:
            run.font.underline = True
        if self.bold:
            run.font.bold = True
        if self.italic:
            run.font.italic = True
        return run

    # ---------------------------------------------------------------- tables

    def _consume_table(self, tokens: Sequence[Token], start: int) -> int:
        """Build a real Word table from ``table_open`` .. ``table_close``."""
        rows: list[list[list[Token]]] = []
        current_row: list[list[Token]] | None = None
        current_cell: list[Token] | None = None
        header_rows = 0
        in_head = False
        index = start + 1
        while index < len(tokens):
            token = tokens[index]
            kind = token.type
            if kind == "table_close":
                index += 1
                break
            if kind == "thead_open":
                in_head = True
            elif kind == "thead_close":
                in_head = False
            elif kind == "tr_open":
                current_row = []
            elif kind == "tr_close":
                if current_row is not None:
                    rows.append(current_row)
                    if in_head:
                        header_rows += 1
                current_row = None
            elif kind in {"th_open", "td_open"}:
                current_cell = []
            elif kind in {"th_close", "td_close"}:
                if current_cell is not None and current_row is not None:
                    current_row.append(current_cell)
                current_cell = None
            elif kind == "inline" and current_cell is not None:
                current_cell.extend(token.children or [])
            index += 1

        if rows:
            self._end_paragraph()
            self._add_table(rows, header_rows=header_rows)
        return index

    def _add_table(self, rows: list[list[list[Token]]], *, header_rows: int) -> None:
        columns = max(len(row) for row in rows)
        table = self.document.add_table(rows=len(rows), cols=columns)
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for column in range(columns):
                cell = table.cell(row_index, column)
                self._end_paragraph()
                self.paragraph = cell.paragraphs[0]
                if column >= len(row):
                    continue
                if row_index < header_rows:
                    self.bold += 1
                for token in row[column]:
                    self._inline_child(token)
                if row_index < header_rows:
                    self.bold -= 1
        self._end_paragraph()


def _force_font(run: Run, name: str) -> None:
    """python-docx's ``font.name`` only sets the ascii/hAnsi slots."""
    properties = run._r.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), name)


def render_transcript_docx(
    *,
    title: str,
    sections: Iterable[TranscriptSection],
    subtitle: str | None = None,
    footer: str | None = None,
    exported_at: datetime | None = None,
) -> bytes:
    """Render a transcript to ``.docx`` bytes.

    ``title`` becomes the document title, each section an author heading with
    its markdown body nested one heading level below, and ``footer`` a closing
    note (e.g. the exported-at line).
    """
    document = Document()
    document.add_heading(title, level=0)
    stamp = (exported_at or datetime.now().astimezone()).strftime("%Y-%m-%d %H:%M %Z").strip()
    caption = (
        f"{subtitle}\nExported from Garbanzo AI · {stamp}"
        if subtitle
        else (f"Exported from Garbanzo AI · {stamp}")
    )
    caption_run = document.add_paragraph().add_run(caption)
    caption_run.font.size = Pt(9)
    caption_run.font.color.rgb = MUTED_COLOR

    writer = _MarkdownWriter(document, heading_offset=1)
    for section in sections:
        document.add_paragraph()
        document.add_heading(section.author, level=1)
        if section.meta:
            meta_run = document.add_paragraph().add_run(section.meta)
            meta_run.font.size = Pt(8)
            meta_run.font.italic = True
            meta_run.font.color.rgb = MUTED_COLOR
        writer.write(section.body)

    if footer:
        document.add_paragraph()
        footer_run = document.add_paragraph().add_run(footer)
        footer_run.font.size = Pt(8)
        footer_run.font.italic = True
        footer_run.font.color.rgb = MUTED_COLOR

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
